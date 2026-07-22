from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "HP_EliteDesk_LocalAI_Remote_Setup_Guide.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(10)


def set_table_borders(table, color="D9E2EC"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_note(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.35)
    cell = table.cell(0, 0)
    cell.width = Inches(6.35)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(31, 77, 120)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    set_table_borders(table, "D9E2EC")
    doc.add_paragraph()


def add_command(doc, command):
    p = doc.add_paragraph()
    p.style = "Command"
    r = p.add_run(command)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)


def add_check(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def make_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    if "Command" not in styles:
        cmd = styles.add_style("Command", 1)
        cmd.font.name = "Consolas"
        cmd.font.size = Pt(9.5)
        cmd.paragraph_format.left_indent = Inches(0.25)
        cmd.paragraph_format.space_before = Pt(2)
        cmd.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("HP EliteDesk Remote Setup Guide for LocalAI Development")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("One-time setup using a temporary monitor, keyboard, and mouse")
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(85, 85, 85)

    add_note(
        doc,
        "Objective",
        "Set up the HP EliteDesk so it can be controlled from the laptop without permanently needing a separate monitor, keyboard, or mouse. The HP will run LocalAI and use the RTX 2080 GPU; the laptop will be used for remote control, coding, browser access, and API testing.",
    )

    doc.add_heading("Target Arrangement", level=1)
    p = doc.add_paragraph()
    p.add_run("Final workflow: ").bold = True
    p.add_run("Laptop -> Remote Desktop or SSH -> HP EliteDesk with RTX 2080 -> LocalAI API available on the network.")

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.9)
    table.columns[1].width = Inches(4.45)
    headers = table.rows[0].cells
    set_cell_text(headers[0], "Item", True)
    set_cell_text(headers[1], "Purpose", True)
    for cell in headers:
        set_cell_shading(cell, "E8EEF5")
    rows = [
        ("Temporary monitor", "Needed only during initial setup."),
        ("Temporary keyboard and mouse", "Needed only to log in and enable remote access."),
        ("HP EliteDesk", "Runs Windows now, later prepared for LocalAI development."),
        ("Laptop", "Used to control the HP after setup."),
        ("Wi-Fi or Ethernet", "Both devices must be on the same network."),
    ]
    for item, purpose in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], item)
        set_cell_text(cells[1], purpose)
    set_table_borders(table)

    doc.add_heading("Phase 1: Physical Setup", level=1)
    add_numbered(doc, "Connect the monitor to the HP. If the HP has a separate RTX 2080 GPU, use the GPU display ports first. These are usually lower on the back panel.")
    add_numbered(doc, "Connect the USB keyboard and USB mouse to the HP. Rear USB ports are preferred.")
    add_numbered(doc, "Connect the HP to the internet. Ethernet is best, but Wi-Fi is acceptable.")
    add_numbered(doc, "Power on the monitor first, then power on the HP EliteDesk.")
    add_numbered(doc, "Log in to Windows using the provided username and password.")

    add_note(
        doc,
        "Important display note",
        "Do not connect HP HDMI directly to the laptop HDMI port. Most laptop HDMI ports are output only. For this setup, the monitor must connect directly to the HP during initial configuration.",
    )

    doc.add_heading("Phase 2: Connect HP to Internet", level=1)
    add_numbered(doc, "Click the Wi-Fi or network icon in the bottom-right taskbar.")
    add_numbered(doc, "Select the correct Wi-Fi network and enter the password.")
    add_numbered(doc, "Open a browser and confirm that https://google.com loads.")

    doc.add_heading("Phase 3: Stop HP From Sleeping", level=1)
    add_numbered(doc, "Open Start and search for Power & sleep settings.")
    add_numbered(doc, "Set Screen to Never if available.")
    add_numbered(doc, "Set Sleep to Never, especially for the plugged-in power option.")
    add_note(doc, "Why this matters", "If the HP sleeps, Remote Desktop and SSH connections from the laptop will stop working.")

    doc.add_heading("Phase 4: Find HP IP Address", level=1)
    add_numbered(doc, "On the HP, press Windows key + R.")
    add_numbered(doc, "Type cmd and press Enter.")
    add_numbered(doc, "Run this command:")
    add_command(doc, "ipconfig")
    p = doc.add_paragraph()
    p.add_run("Find the active adapter, usually ")
    p.add_run("Wireless LAN adapter Wi-Fi").bold = True
    p.add_run(" or ")
    p.add_run("Ethernet adapter Ethernet").bold = True
    p.add_run(". Write down the IPv4 Address, for example 192.168.1.50.")

    doc.add_heading("Phase 5: Find HP Username", level=1)
    add_numbered(doc, "In the same Command Prompt, run:")
    add_command(doc, "whoami")
    p = doc.add_paragraph()
    p.add_run("Example output: ").bold = True
    p.add_run("hpdesk\\ali")
    p = doc.add_paragraph()
    p.add_run("The username is the part after the slash: ").bold = True
    p.add_run("ali")

    doc.add_heading("Phase 6: Check Windows Edition", level=1)
    add_numbered(doc, "On the HP, press Windows key + R.")
    add_numbered(doc, "Type winver and press Enter.")
    add_numbered(doc, "Check if it says Windows Pro, Enterprise, or Home.")
    add_note(
        doc,
        "Remote Desktop rule",
        "Windows Pro and Enterprise can usually host Remote Desktop. Windows Home usually cannot host Remote Desktop, so SSH should be used instead.",
    )

    doc.add_heading("Phase 7A: Enable Remote Desktop if Windows Pro or Enterprise", level=1)
    add_numbered(doc, "Open Start and search for Remote Desktop settings.")
    add_numbered(doc, "Turn on Enable Remote Desktop.")
    add_numbered(doc, "Confirm the prompt if Windows asks.")
    add_numbered(doc, "On the laptop, open Remote Desktop Connection.")
    add_numbered(doc, "Enter the HP IP address, then connect using the HP username and password.")
    add_note(doc, "Success condition", "If the HP Windows desktop appears on the laptop, Remote Desktop is working.")

    doc.add_heading("Phase 7B: Enable SSH on Windows HP", level=1)
    add_numbered(doc, "On HP Windows, open Optional features from the Start menu.")
    add_numbered(doc, "Click View features or Add a feature.")
    add_numbered(doc, "Search for OpenSSH Server, select it, and install it.")
    add_numbered(doc, "Open PowerShell as Administrator.")
    add_numbered(doc, "Run these commands one by one:")
    add_command(doc, "Start-Service sshd")
    add_command(doc, "Set-Service -Name sshd -StartupType Automatic")
    add_command(doc, "Get-Service sshd")
    p = doc.add_paragraph()
    p.add_run("Expected result: ").bold = True
    p.add_run("Status should be Running.")
    add_numbered(doc, "Allow SSH through Windows Firewall:")
    add_command(doc, 'New-NetFirewallRule -Name sshd -DisplayName "OpenSSH Server" -Enabled True -Direction Inbound -Protocol TCP -Action Allow -LocalPort 22')

    doc.add_heading("Phase 8: Test SSH From Laptop", level=1)
    add_numbered(doc, "On the laptop, open PowerShell.")
    add_numbered(doc, "Run this command, replacing the username and IP address:")
    add_command(doc, "ssh HP_USERNAME@HP_IP")
    p = doc.add_paragraph()
    p.add_run("Example: ").bold = True
    p.add_run("ssh ali@192.168.1.50")
    add_numbered(doc, "If asked whether to continue connecting, type yes.")
    add_numbered(doc, "Enter the HP Windows password.")
    add_note(doc, "Success condition", "If the laptop shows a prompt like C:\\Users\\ali>, SSH is working.")

    doc.add_heading("Phase 9: Check RTX 2080 GPU", level=1)
    add_numbered(doc, "On the HP, open Command Prompt or PowerShell.")
    add_numbered(doc, "Run:")
    add_command(doc, "nvidia-smi")
    p = doc.add_paragraph()
    p.add_run("Expected result: ").bold = True
    p.add_run("The output should show NVIDIA GeForce RTX 2080, driver version, CUDA version, and GPU memory.")
    add_numbered(doc, "If nvidia-smi is not recognized, run:")
    add_command(doc, 'cd "C:\\Program Files\\NVIDIA Corporation\\NVSMI"')
    add_command(doc, "nvidia-smi")
    add_note(doc, "If this still fails", "The NVIDIA driver may need to be installed or repaired. Do not continue into LocalAI GPU setup until the GPU is visible.")

    doc.add_heading("Phase 10: Disconnect Temporary Devices", level=1)
    p = doc.add_paragraph()
    p.add_run("Only disconnect the borrowed monitor, keyboard, and mouse after confirming at least one remote method works from the laptop:")
    add_check(doc, "Remote Desktop works from the laptop.")
    add_check(doc, "SSH works from the laptop.")
    p = doc.add_paragraph()
    p.add_run("Keep connected: ").bold = True
    p.add_run("HP power and HP network connection.")

    doc.add_heading("Information to Send Back", level=1)
    items = [
        "Windows edition: Home / Pro / Enterprise",
        "HP IP address",
        "HP username",
        "Remote Desktop working: yes / no / not available",
        "SSH working: yes / no",
        "nvidia-smi shows RTX 2080: yes / no",
        "HP internet connection: Wi-Fi / Ethernet",
    ]
    for item in items:
        add_check(doc, item)

    doc.add_heading("Notes for LocalAI Suitability", level=1)
    p = doc.add_paragraph()
    p.add_run("The RTX 2080 is suitable for LocalAI development, UI changes, API testing, and small to medium quantized models. ")
    p.add_run("Start with 1B to 4B quantized models, then test a 7B Q4 model if needed. Avoid large 30B or 70B models on this GPU.")

    doc.add_heading("Recommended Next Step After Remote Access Works", level=1)
    p = doc.add_paragraph()
    p.add_run("After this setup is complete, choose the development environment. For LocalAI and NVIDIA GPU work, the clean long-term options are:")
    add_check(doc, "Windows plus WSL2 Ubuntu")
    add_check(doc, "Full Ubuntu installation on the HP")
    add_check(doc, "Temporary Docker-based setup while learning the project")

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("HP EliteDesk LocalAI Setup Guide")
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(85, 85, 85)

    doc.save(OUT)


if __name__ == "__main__":
    make_doc()
